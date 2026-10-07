# Backend tests for /api/contingency/export (NEW SPEC: all applications left-join)
# and regressions.
import io
import os
import re

import openpyxl
import pytest
import requests

BASE_URL = os.environ["EXPO_PUBLIC_BACKEND_URL"].rstrip("/")
API = f"{BASE_URL}/api"

EXPECTED_HEADERS = [
    "Platform",
    "Business Application",
    "Business Process",
    "Contingency Plan (Failed System)",
    "Owner",
    "Has Contingency Plan",
]
EXPECTED_PRODUCT_COUNT = 150          # from seed_data.PRODUCTS
EXPECTED_YES_ROWS = 37                # from seed_data.CONTINGENCY_MAP
EXPECTED_NO_ROWS = 122                # 150 products - 28 distinct product names in CONTINGENCY_MAP = 122
EXPECTED_TOTAL_DATA_ROWS = EXPECTED_YES_ROWS + EXPECTED_NO_ROWS  # 159


@pytest.fixture(scope="module")
def session():
    return requests.Session()


# --- /api/contingency/export (NEW SPEC) ----------------------------------
class TestContingencyExport:
    def test_export_status_and_headers(self, session):
        r = session.get(f"{API}/contingency/export", timeout=30)
        assert r.status_code == 200, r.text
        ct = r.headers.get("content-type", "")
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in ct, ct
        cd = r.headers.get("content-disposition", "")
        assert "attachment" in cd.lower(), cd
        m = re.search(r'filename="?(contingency_mapping_\d{8}\.xlsx)"?', cd)
        assert m, f"Content-Disposition missing dated filename: {cd!r}"

    def test_xlsx_sheet_header_and_freeze(self, session):
        r = session.get(f"{API}/contingency/export", timeout=30)
        assert r.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
        assert "Contingency Mapping" in wb.sheetnames
        ws = wb["Contingency Mapping"]

        header = [c.value for c in ws[1]]
        assert header == EXPECTED_HEADERS, f"Header mismatch: {header}"

        assert ws.freeze_panes == "A2", f"freeze_panes={ws.freeze_panes!r}"

        # Red header fill E61A27
        first_cell = ws["A1"]
        fg = (
            first_cell.fill.fgColor.rgb
            if first_cell.fill and first_cell.fill.fgColor
            else ""
        )
        assert fg and fg.upper().endswith("E61A27"), f"Header fill color unexpected: {fg}"

    def test_xlsx_row_coverage(self, session):
        r = session.get(f"{API}/contingency/export", timeout=30)
        wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
        ws = wb["Contingency Mapping"]

        data_rows = [
            [c.value for c in row]
            for row in ws.iter_rows(min_row=2, values_only=False)
            if any(c.value not in (None, "") for c in row)
        ]

        # Reference count from live /api/products (ground truth)
        products = session.get(f"{API}/products", timeout=15).json()
        product_count = len(products)
        assert product_count == EXPECTED_PRODUCT_COUNT, (
            f"Product catalog has {product_count}, expected {EXPECTED_PRODUCT_COUNT}"
        )

        # Distinct Business Application values must cover all 150 products
        distinct_apps = {str(r[1]) for r in data_rows if r[1]}
        assert len(distinct_apps) == EXPECTED_PRODUCT_COUNT, (
            f"Distinct apps in export: {len(distinct_apps)}, expected {EXPECTED_PRODUCT_COUNT}"
        )

        yes_rows = [r for r in data_rows if r[5] == "Yes"]
        no_rows = [r for r in data_rows if r[5] == "No"]

        assert len(yes_rows) == EXPECTED_YES_ROWS, (
            f"'Yes' rows: {len(yes_rows)}, expected {EXPECTED_YES_ROWS}"
        )
        assert len(no_rows) == EXPECTED_NO_ROWS, (
            f"'No' rows: {len(no_rows)}, expected {EXPECTED_NO_ROWS}"
        )
        assert len(data_rows) == EXPECTED_TOTAL_DATA_ROWS, (
            f"Total data rows: {len(data_rows)}, expected {EXPECTED_TOTAL_DATA_ROWS}"
        )

    def test_yes_rows_populated_no_rows_blank(self, session):
        r = session.get(f"{API}/contingency/export", timeout=30)
        wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
        ws = wb["Contingency Mapping"]
        data_rows = [
            [c.value for c in row]
            for row in ws.iter_rows(min_row=2, values_only=False)
            if any(c.value not in (None, "") for c in row)
        ]

        for row in data_rows:
            platform, app, process, plan, owner, flag = row
            assert app not in (None, ""), f"Business Application blank in row: {row}"
            if flag == "Yes":
                assert process and str(process).strip(), f"Yes row w/ blank process: {row}"
                assert owner and str(owner).strip(), f"Yes row w/ blank owner: {row}"
            elif flag == "No":
                assert (process in (None, "")), f"No row w/ non-blank process: {row}"
                assert (plan in (None, "")), f"No row w/ non-blank plan: {row}"
                assert (owner in (None, "")), f"No row w/ non-blank owner: {row}"
            else:
                pytest.fail(f"Unexpected 'Has Contingency Plan' value: {flag!r}")


# --- Regressions ---------------------------------------------------------
class TestRegressions:
    def test_products_count_is_150(self, session):
        r = session.get(f"{API}/products", timeout=15)
        assert r.status_code == 200
        data = r.json()
        assert isinstance(data, list)
        assert len(data) == EXPECTED_PRODUCT_COUNT, (
            f"/api/products returned {len(data)}, expected {EXPECTED_PRODUCT_COUNT}"
        )

    def test_contingency_sfa_returns_two_rows(self, session):
        r = session.get(f"{API}/contingency", params={"product": "SFA"}, timeout=15)
        assert r.status_code == 200
        data = r.json()
        assert isinstance(data, list)
        assert len(data) == 2, f"Expected 2 rows for SFA, got {len(data)}: {data}"
        for row in data:
            assert row["product"] == "SFA"
            assert "failed_system" in row and "owner" in row and "process" in row

    def test_contingency_empty_query_returns_empty(self, session):
        r = session.get(f"{API}/contingency", timeout=15)
        assert r.status_code == 200
        assert r.json() == []

    def test_catalog_endpoints_up(self, session):
        for ep in ("/products", "/business-processes", "/countries", "/templates"):
            r = session.get(f"{API}{ep}", timeout=15)
            assert r.status_code == 200, f"{ep} -> {r.status_code}"
            assert isinstance(r.json(), list)

    def test_render_html_countries_newline(self, session):
        tpl = session.get(f"{API}/templates", timeout=15).json()
        imcr = next((t for t in tpl if t.get("category") == "IMCR"), None)
        assert imcr, "No IMCR template"
        d = session.post(f"{API}/drafts", json={"template_id": imcr["id"]}, timeout=15).json()
        draft_id = d["id"]
        try:
            vals = dict(d.get("values", {}))
            vals["subject"] = "TEST_render_countries"
            vals["countries"] = '["Austria", "Italy", "Poland"]'
            session.put(
                f"{API}/drafts/{draft_id}",
                json={"values": vals},
                timeout=15,
            )
            r = session.get(f"{API}/drafts/{draft_id}/render-html", timeout=15)
            assert r.status_code == 200
            html = r.json()["html"]
            assert "Austria<br>Italy<br>Poland" in html
        finally:
            session.delete(f"{API}/drafts/{draft_id}", timeout=15)
