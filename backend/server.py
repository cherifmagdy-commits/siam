import io
import logging
import os
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
import openpyxl.styles
import openpyxl.utils
import httpx
from dotenv import load_dotenv
from fastapi import APIRouter, FastAPI, HTTPException, UploadFile, File, Header
from fastapi.responses import StreamingResponse
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field
from starlette.middleware.cors import CORSMiddleware

from seed_data import (
    BUSINESS_PROCESSES,
    CENTRAL_DTPS,
    CONTINGENCY_MAP,
    COUNTRIES,
    PLATFORM_DIRECTORS_CENTRAL,
    PRODUCTS,
)
from templates_seed import STAGE_LABEL, STAGES, build_templates
from outlook import register_outlook_routes, create_shared_draft as _create_shared_draft

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
)
logger = logging.getLogger(__name__)


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def new_id() -> str:
    return str(uuid.uuid4())


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
class ProductCreate(BaseModel):
    name: str
    platform: str = "Custom"


class ProcessCreate(BaseModel):
    name: str


class CountryCreate(BaseModel):
    name: str
    cluster: str = "Custom"
    country_dtps: List[str] = []
    country_dl: List[str] = []
    dwt_leader: List[str] = []
    platform_directors: List[str] = []


class DraftCreate(BaseModel):
    template_id: str


class IncidentCreate(BaseModel):
    template_id: str


class UpdateCreate(BaseModel):
    template_id: str


class MarkSent(BaseModel):
    draft_id: str


class RegisterPushBody(BaseModel):
    user_id: str
    platform: str
    device_token: str


class DraftUpdate(BaseModel):
    values: Optional[Dict[str, Any]] = None
    validations: Optional[Dict[str, bool]] = None
    recipients: Optional[str] = None


# ---------------------------------------------------------------------------
# Seeding
# ---------------------------------------------------------------------------
async def seed_database():
    if await db.products.count_documents({}) == 0:
        docs = [
            {"id": new_id(), "name": p["name"], "platform": p["platform"],
             "custom": False, "deleted_at": None, "created_at": now_iso()}
            for p in PRODUCTS
        ]
        if docs:
            await db.products.insert_many(docs)
        logger.info("Seeded %d products", len(docs))

    if await db.business_processes.count_documents({}) == 0:
        docs = [
            {"id": new_id(), "name": name, "custom": False,
             "deleted_at": None, "created_at": now_iso()}
            for name in BUSINESS_PROCESSES
        ]
        if docs:
            await db.business_processes.insert_many(docs)
        logger.info("Seeded %d business processes", len(docs))

    if await db.contingency_map.count_documents({}) == 0:
        docs = [
            {"id": new_id(), "product": m["product"], "process": m["process"],
             "failed_system": m["failed_system"], "owner": m["owner"]}
            for m in CONTINGENCY_MAP
        ]
        if docs:
            await db.contingency_map.insert_many(docs)
        logger.info("Seeded %d contingency map rows", len(docs))

    if await db.countries.count_documents({}) == 0:
        docs = [
            {"id": new_id(), "name": c["name"], "cluster": c["cluster"],
             "country_dtps": c["country_dtps"], "country_dl": c["country_dl"],
             "dwt_leader": c["dwt_leader"], "platform_directors": c["platform_directors"],
             "custom": False, "deleted_at": None, "created_at": now_iso()}
            for c in COUNTRIES
        ]
        if docs:
            await db.countries.insert_many(docs)
        logger.info("Seeded %d countries", len(docs))

    # Templates: reseed every startup so structure changes propagate.
    templates = build_templates()
    for t in templates:
        t["id"] = t["key"]
        await db.templates.replace_one({"key": t["key"]}, t, upsert=True)
    logger.info("Seeded %d templates", len(templates))


@app.on_event("startup")
async def on_startup():
    await seed_database()


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()


# ---------------------------------------------------------------------------
# Products
# ---------------------------------------------------------------------------
@api_router.get("/products")
async def list_products():
    cursor = db.products.find({"deleted_at": None}, {"_id": 0}).sort("name", 1)
    items = await cursor.to_list(2000)
    return items


@api_router.post("/products")
async def create_product(payload: ProductCreate):
    name = payload.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Name is required")
    existing = await db.products.find_one({"name": name, "deleted_at": None})
    if existing:
        existing.pop("_id", None)
        return existing
    doc = {"id": new_id(), "name": name, "platform": payload.platform.strip() or "Custom",
           "custom": True, "deleted_at": None, "created_at": now_iso()}
    await db.products.insert_one(dict(doc))
    return doc


@api_router.delete("/products/{item_id}")
async def delete_product(item_id: str):
    res = await db.products.update_one({"id": item_id}, {"$set": {"deleted_at": now_iso()}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Product not found")
    return {"ok": True}


# ---------------------------------------------------------------------------
# Business processes
# ---------------------------------------------------------------------------
@api_router.get("/business-processes")
async def list_processes():
    cursor = db.business_processes.find({"deleted_at": None}, {"_id": 0}).sort("name", 1)
    return await cursor.to_list(2000)


@api_router.post("/business-processes")
async def create_process(payload: ProcessCreate):
    name = payload.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Name is required")
    existing = await db.business_processes.find_one({"name": name, "deleted_at": None})
    if existing:
        existing.pop("_id", None)
        return existing
    doc = {"id": new_id(), "name": name, "custom": True,
           "deleted_at": None, "created_at": now_iso()}
    await db.business_processes.insert_one(dict(doc))
    return doc


@api_router.delete("/business-processes/{item_id}")
async def delete_process(item_id: str):
    res = await db.business_processes.update_one({"id": item_id}, {"$set": {"deleted_at": now_iso()}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Process not found")
    return {"ok": True}


# ---------------------------------------------------------------------------
# Countries (crisis distribution list)
# ---------------------------------------------------------------------------
@api_router.get("/countries")
async def list_countries():
    cursor = db.countries.find({"deleted_at": None}, {"_id": 0}).sort("name", 1)
    return await cursor.to_list(500)


@api_router.post("/countries")
async def create_country(payload: CountryCreate):
    name = payload.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Name is required")
    existing = await db.countries.find_one({"name": name, "deleted_at": None})
    if existing:
        existing.pop("_id", None)
        return existing
    doc = {
        "id": new_id(), "name": name, "cluster": payload.cluster.strip() or "Custom",
        "country_dtps": payload.country_dtps, "country_dl": payload.country_dl,
        "dwt_leader": payload.dwt_leader, "platform_directors": payload.platform_directors,
        "custom": True, "deleted_at": None, "created_at": now_iso(),
    }
    await db.countries.insert_one(dict(doc))
    return doc


@api_router.delete("/countries/{item_id}")
async def delete_country(item_id: str):
    res = await db.countries.update_one({"id": item_id}, {"$set": {"deleted_at": now_iso()}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Country not found")
    return {"ok": True}


# ---------------------------------------------------------------------------
# Templates
# ---------------------------------------------------------------------------
@api_router.get("/templates")
async def list_templates(category: Optional[str] = None):
    query: Dict[str, Any] = {}
    if category:
        query["category"] = category
    cursor = db.templates.find(query, {"_id": 0})
    return await cursor.to_list(100)


@api_router.get("/templates/{template_id}")
async def get_template(template_id: str):
    doc = await db.templates.find_one({"id": template_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Template not found")
    return doc


# ---------------------------------------------------------------------------
# Drafts
# ---------------------------------------------------------------------------
def _default_values(template: Dict[str, Any]) -> Dict[str, Any]:
    values: Dict[str, Any] = {}
    for section in template["sections"]:
        for field in section["fields"]:
            values[field["key"]] = field.get("default", "")
    return values


@api_router.post("/drafts")
async def create_draft(payload: DraftCreate):
    template = await db.templates.find_one({"id": payload.template_id}, {"_id": 0})
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    doc = {
        "id": new_id(),
        "template_id": template["id"],
        "template_name": template["name"],
        "category": template["category"],
        "stage": template["stage"],
        "values": _default_values(template),
        "validations": {},
        "recipients": "",
        "deleted_at": None,
        "created_at": now_iso(),
        "updated_at": now_iso(),
    }
    await db.drafts.insert_one(dict(doc))
    return doc


@api_router.get("/drafts")
async def list_drafts():
    cursor = db.drafts.find({"deleted_at": None}, {"_id": 0}).sort("updated_at", -1)
    drafts = await cursor.to_list(500)
    # attach validation progress
    for d in drafts:
        template = await db.templates.find_one({"id": d["template_id"]}, {"_id": 0})
        total = len(template["sections"]) if template else 0
        validated = sum(1 for v in d.get("validations", {}).values() if v)
        d["progress"] = {"validated": validated, "total": total}
    return drafts


@api_router.get("/drafts/{draft_id}")
async def get_draft(draft_id: str):
    doc = await db.drafts.find_one({"id": draft_id, "deleted_at": None}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Draft not found")
    template = await db.templates.find_one({"id": doc["template_id"]}, {"_id": 0})
    doc["template"] = template
    return doc


@api_router.put("/drafts/{draft_id}")
async def update_draft(draft_id: str, payload: DraftUpdate):
    doc = await db.drafts.find_one({"id": draft_id, "deleted_at": None})
    if not doc:
        raise HTTPException(status_code=404, detail="Draft not found")
    update: Dict[str, Any] = {"updated_at": now_iso()}
    if payload.values is not None:
        update["values"] = payload.values
    if payload.validations is not None:
        update["validations"] = payload.validations
    if payload.recipients is not None:
        update["recipients"] = payload.recipients
    await db.drafts.update_one({"id": draft_id}, {"$set": update})
    updated = await db.drafts.find_one({"id": draft_id}, {"_id": 0})
    return updated


@api_router.delete("/drafts/{draft_id}")
async def delete_draft(draft_id: str):
    res = await db.drafts.update_one({"id": draft_id}, {"$set": {"deleted_at": now_iso()}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Draft not found")
    return {"ok": True}


# ---------------------------------------------------------------------------
# Incidents (a linked thread of updates through the stages)
# ---------------------------------------------------------------------------
def _new_draft_doc(template, incident_id, sequence, prefill):
    values = _default_values(template)
    if prefill:
        for k in list(values.keys()):
            pv = prefill.get(k)
            if pv not in (None, ""):
                values[k] = pv
        for extra in ("product_team", "sre"):
            if prefill.get(extra):
                values[extra] = prefill[extra]
    return {
        "id": new_id(),
        "template_id": template["id"],
        "template_name": template["name"],
        "category": template["category"],
        "stage": template["stage"],
        "incident_id": incident_id,
        "sequence": sequence,
        "values": values,
        "validations": {},
        "recipients": "",
        "deleted_at": None,
        "created_at": now_iso(),
        "updated_at": now_iso(),
    }


def _incident_title(values, template):
    return str(values.get("title") or values.get("subject") or template["name"]).strip()


@api_router.post("/incidents")
async def create_incident(payload: IncidentCreate):
    template = await db.templates.find_one({"id": payload.template_id}, {"_id": 0})
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    incident_id = new_id()
    draft = _new_draft_doc(template, incident_id, 1, None)
    await db.drafts.insert_one(dict(draft))
    inc = {
        "id": incident_id,
        "title": _incident_title(draft["values"], template),
        "category": template["category"],
        "current_stage": template["stage"],
        "status": "OPEN",
        "last_sent_at": None,
        "next_update_at": None,
        "created_at": now_iso(),
        "updated_at": now_iso(),
        "deleted_at": None,
    }
    await db.incidents.insert_one(dict(inc))
    return {"incident": inc, "draft": draft}


@api_router.post("/incidents/{incident_id}/updates")
async def create_update(incident_id: str, payload: UpdateCreate):
    inc = await db.incidents.find_one({"id": incident_id, "deleted_at": None})
    if not inc:
        raise HTTPException(status_code=404, detail="Incident not found")
    template = await db.templates.find_one({"id": payload.template_id}, {"_id": 0})
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    latest = await db.drafts.find(
        {"incident_id": incident_id, "deleted_at": None}
    ).sort("sequence", -1).to_list(1)
    prefill = latest[0]["values"] if latest else None
    seq = (latest[0]["sequence"] + 1) if latest else 1
    draft = _new_draft_doc(template, incident_id, seq, prefill)
    await db.drafts.insert_one(dict(draft))
    await db.incidents.update_one({"id": incident_id}, {"$set": {"updated_at": now_iso()}})
    return draft


@api_router.get("/incidents")
async def list_incidents():
    cursor = db.incidents.find({"deleted_at": None}, {"_id": 0}).sort("updated_at", -1)
    incidents = await cursor.to_list(500)
    for inc in incidents:
        inc["update_count"] = await db.drafts.count_documents(
            {"incident_id": inc["id"], "deleted_at": None}
        )
    return incidents


@api_router.get("/incidents/{incident_id}")
async def get_incident(incident_id: str):
    inc = await db.incidents.find_one({"id": incident_id, "deleted_at": None}, {"_id": 0})
    if not inc:
        raise HTTPException(status_code=404, detail="Incident not found")
    cursor = db.drafts.find({"incident_id": incident_id, "deleted_at": None}, {"_id": 0}).sort("sequence", 1)
    drafts = await cursor.to_list(200)
    for d in drafts:
        template = await db.templates.find_one({"id": d["template_id"]}, {"_id": 0})
        total = len(template["sections"]) if template else 0
        validated = sum(1 for v in d.get("validations", {}).values() if v)
        d["progress"] = {"validated": validated, "total": total}
    inc["updates"] = drafts
    return inc


@api_router.post("/incidents/{incident_id}/mark-sent")
async def mark_sent(incident_id: str, payload: MarkSent):
    inc = await db.incidents.find_one({"id": incident_id, "deleted_at": None})
    if not inc:
        raise HTTPException(status_code=404, detail="Incident not found")
    draft = await db.drafts.find_one({"id": payload.draft_id, "deleted_at": None}, {"_id": 0})
    if not draft:
        raise HTTPException(status_code=404, detail="Draft not found")
    update: Dict[str, Any] = {
        "last_sent_at": now_iso(),
        "current_stage": draft["stage"],
        "updated_at": now_iso(),
    }
    if draft["stage"] == "RESOLVED":
        update["status"] = "RESOLVED"
        update["resolved_at"] = now_iso()
    await db.incidents.update_one({"id": incident_id}, {"$set": update})
    await db.drafts.update_one({"id": payload.draft_id}, {"$set": {"sent_at": now_iso()}})
    updated = await db.incidents.find_one({"id": incident_id}, {"_id": 0})
    return updated


@api_router.delete("/incidents/{incident_id}")
async def delete_incident(incident_id: str):
    res = await db.incidents.update_one({"id": incident_id}, {"$set": {"deleted_at": now_iso()}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Incident not found")
    await db.drafts.update_many({"incident_id": incident_id}, {"$set": {"deleted_at": now_iso()}})
    return {"ok": True}


# ---------------------------------------------------------------------------
# Push notifications (Emergent managed relay)
# ---------------------------------------------------------------------------
PUSH_BASE_URL = "https://integrations.emergentagent.com"
PUSH_KEY = os.environ.get("EMERGENT_PUSH_KEY", "placeholder")
_push_client = httpx.AsyncClient(base_url=PUSH_BASE_URL, headers={"X-Push-Key": PUSH_KEY}, timeout=10.0)


@api_router.post("/register-push", status_code=201)
async def register_push(body: RegisterPushBody):
    resp = await _push_client.post("/api/v1/push/users/register", json=body.model_dump())
    if resp.status_code == 401:
        raise HTTPException(500, "EMERGENT_PUSH_KEY missing or invalid")
    if resp.status_code >= 500:
        raise HTTPException(502, "Push provider unavailable")
    resp.raise_for_status()
    return {"status": "registered"}


async def send_push(recipients: List[str], data: Dict[str, Any], idempotency_key: Optional[str] = None) -> None:
    if not recipients:
        return
    if "title" not in data or "message" not in data:
        raise ValueError("data must include title and message")
    payload: Dict[str, Any] = {"recipients": recipients[:100], "data": data}
    if idempotency_key:
        payload["$idempotency_key"] = idempotency_key
    resp = await _push_client.post("/api/v1/push/trigger", json=payload)
    if resp.status_code == 401:
        raise HTTPException(500, "EMERGENT_PUSH_KEY missing or invalid")
    if resp.status_code >= 500:
        raise HTTPException(502, "Push provider unavailable")
    resp.raise_for_status()


import json as _json


def _format_countries(raw: str) -> str:
    try:
        names = _json.loads(raw)
        if isinstance(names, list):
            return "\n".join(str(n) for n in names)
    except Exception:
        pass
    return raw


def _render_body(template: Dict[str, Any], values: Dict[str, Any]) -> str:
    lines: List[str] = []
    lines.append(template.get("header", ""))
    lines.append("")
    for section in template["sections"]:
        if section["key"] == "subject":
            continue
        section_lines: List[str] = []
        single_field = len(section["fields"]) == 1
        for field in section["fields"]:
            raw = values.get(field["key"], "")
            if field["type"] == "countries":
                val = _format_countries(str(raw or "")).strip()
            else:
                val = str(raw or "").strip()
            if not val:
                continue
            if single_field and field["type"] == "textarea":
                section_lines.append(val)
            else:
                section_lines.append(f"{field['label']}: {val}")
        if section_lines:
            lines.append(f"── {section['title'].upper()} ──")
            lines.extend(section_lines)
            lines.append("")
    lines.append("—")
    lines.append(template.get("footer", ""))
    return "\n".join(lines).strip()


def _split_manual(raw: str) -> List[str]:
    if not raw:
        return []
    parts = re.split(r"[;,\s]+", raw.strip())
    return [p for p in parts if "@" in p]


async def _compute_recipients(doc: Dict[str, Any]) -> List[str]:
    values = doc.get("values", {})
    category = doc.get("category")
    recipients: List[str] = [CENTRAL_DTPS, PLATFORM_DIRECTORS_CENTRAL]

    selected: List[str] = []
    try:
        selected = _json.loads(values.get("countries", "") or "[]")
        if not isinstance(selected, list):
            selected = []
    except Exception:
        selected = []

    if selected:
        cursor = db.countries.find({"name": {"$in": selected}, "deleted_at": None}, {"_id": 0})
        countries = await cursor.to_list(200)
        for c in countries:
            recipients.extend(c.get("platform_directors", []))
            if category == "IMCR":
                recipients.extend(c.get("country_dl", []))
                recipients.extend(c.get("country_dtps", []))
                recipients.extend(c.get("dwt_leader", []))

    recipients.extend(_split_manual(str(values.get("product_team", "") or "")))
    recipients.extend(_split_manual(str(values.get("sre", "") or "")))

    # dedupe, case-insensitive, preserve order
    out: List[str] = []
    seen = set()
    for r in recipients:
        r = r.strip()
        if r and r.lower() not in seen:
            seen.add(r.lower())
            out.append(r)
    return out


@api_router.get("/drafts/{draft_id}/render")
async def render_draft(draft_id: str):
    doc = await db.drafts.find_one({"id": draft_id, "deleted_at": None}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Draft not found")
    template = await db.templates.find_one({"id": doc["template_id"]}, {"_id": 0})
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    values = doc.get("values", {})
    subject = str(values.get("subject", "") or doc.get("template_name", "")).strip()
    body = _render_body(template, values)
    to = await _compute_recipients(doc)
    return {"subject": subject, "body": body, "to": to, "recipients": "; ".join(to)}


# ---------------------------------------------------------------------------
# HTML render — mirrors the corporate template layout / fonts / colours
# ---------------------------------------------------------------------------
FONT_STACK = "Aptos, Calibri, Helvetica, Arial, sans-serif"
STAGE_COLOR = {
    "IDENTIFIED": "#E29D6E",
    "INVESTIGATING": "#E6A23C",
    "RECOVERING": "#4FA6C7",
    "MONITORING": "#7C8CA0",
    "RESOLVED": "#4CAF7D",
}


def _esc(v: Any) -> str:
    s = str(v or "")
    return (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;").replace("\n", "<br>"))


def _val(values, key):
    return str(values.get(key, "") or "").strip()


def _field_label(template, key):
    for s in template["sections"]:
        for f in s["fields"]:
            if f["key"] == key:
                return f["label"]
    return key


def _box(label, value, border, extra=""):
    return (
        f'<td style="border:1px solid {border};border-radius:12px;padding:10px 18px;'
        f'vertical-align:top;{extra}">'
        f'<div style="font-size:11px;font-weight:700;letter-spacing:.5px;color:#8a8a8a;'
        f'text-transform:uppercase">{_esc(label)}</div>'
        f'<div style="font-size:14px;font-weight:600;color:#1f2937;margin-top:4px">{_esc(value)}</div>'
        f'</td>'
    )


def _impact_row(label, value):
    if not value:
        return ""
    return (
        f'<tr>'
        f'<td style="padding:9px 12px;border-bottom:1px solid #e8ecf0;font-size:13px;'
        f'color:#6b7280;width:42%;vertical-align:top">{_esc(label)}</td>'
        f'<td style="padding:9px 12px;border-bottom:1px solid #e8ecf0;font-size:13px;'
        f'font-weight:600;color:#1f2937">{_esc(value)}</td>'
        f'</tr>'
    )


def _render_html(template: Dict[str, Any], values: Dict[str, Any]) -> str:
    category = template["category"]
    stage = template["stage"]
    accent = STAGE_COLOR.get(stage, "#E29D6E")
    subject = _val(values, "subject") or template["name"]
    parts: List[str] = []

    # Header band
    parts.append(
        f'<div style="border-bottom:3px solid {accent};padding-bottom:12px;margin-bottom:18px">'
        f'<div style="font-size:12px;font-weight:700;letter-spacing:.6px;color:#6b7280;'
        f'text-transform:uppercase">{_esc(template.get("header",""))}</div>'
        f'<div style="font-size:20px;font-weight:700;color:#111827;margin-top:6px">{_esc(subject)}</div>'
        f'</div>'
    )

    # Status pipeline (IMCR only)
    if category == "IMCR":
        cells = ""
        for st in STAGES:
            active = st == stage
            bg = accent if active else "#E8ECF0"
            color = "#ffffff" if active else "#5b6b7b"
            weight = "700" if active else "500"
            cells += (
                f'<td style="background-color:{bg};color:{color};font-weight:{weight};'
                f'font-size:12px;text-align:center;padding:9px 4px;border-right:2px solid #ffffff">'
                f'{_esc(STAGE_LABEL[st])}</td>'
            )
        parts.append(
            f'<table style="width:100%;border-collapse:collapse;border-radius:8px;'
            f'overflow:hidden;margin-bottom:16px"><tr>{cells}</tr></table>'
        )

    # Status headline band
    headline = _val(values, "headline")
    if headline:
        parts.append(
            f'<div style="background-color:{accent};color:#ffffff;font-size:16px;font-weight:600;'
            f'padding:12px 16px;border-radius:10px;margin-bottom:16px">{_esc(headline)}</div>'
        )

    # Timing boxes
    started = _val(values, "issue_started")
    second_key = "resolved_at" if "resolved_at" in values else ("restored_at" if "restored_at" in values else "next_update")
    second = _val(values, second_key)
    if started or second:
        cells = ""
        if started:
            cells += _box(_field_label(template, "issue_started"), started, "#8E0B0B")
        if started and second:
            cells += '<td style="width:14px"></td>'
        if second:
            cells += _box(_field_label(template, second_key), second, "#8a8a8a")
        parts.append(f'<table style="width:100%;margin-bottom:16px"><tr>{cells}</tr></table>')

    # Advisory note (non-IMCR)
    note = _val(values, "note")
    if note:
        parts.append(
            f'<div style="background-color:#FCFBF7;border:1px solid #FDF2D7;border-radius:10px;'
            f'padding:12px 16px;margin-bottom:16px;font-size:13px;color:#6b5b3a;font-style:italic">{_esc(note)}</div>'
        )

    # Description / resolution
    body = _val(values, "body")
    no_action = _val(values, "no_action")
    if body:
        parts.append(
            f'<div style="background-color:rgba(47,62,80,0.03);border-radius:10px;padding:14px 16px;'
            f'margin-bottom:16px;font-size:14px;line-height:1.5;color:#1f2937">{_esc(body)}'
            + (f'<div style="margin-top:8px;font-weight:600">{_esc(no_action)}</div>' if no_action else "")
            + '</div>'
        )

    # Business impact / details table
    impact_keys = [
        ("business_application", "affected_service"),
        ("business_process",),
        ("countries",),
        ("incident_ref",),
        ("workaround",),
    ]
    rows = ""
    for keys in impact_keys:
        for k in keys:
            if k in values:
                raw = values.get(k, "")
                v = _format_countries(str(raw or "")) if k == "countries" else str(raw or "").strip()
                rows += _impact_row(_field_label(template, k), v)
                break
    if rows:
        parts.append(
            f'<div style="font-size:12px;font-weight:700;letter-spacing:.6px;color:#8a8a8a;'
            f'text-transform:uppercase;margin-bottom:8px">Business Impact</div>'
            f'<table style="width:100%;border-collapse:collapse;border:1px solid #e8ecf0;'
            f'border-radius:10px;overflow:hidden;margin-bottom:16px">{rows}</table>'
        )

    # Contingency box
    contingency = _val(values, "contingency")
    if contingency:
        parts.append(
            f'<div style="background-color:#FCFBF7;border:1px solid #FDF2D7;border-radius:10px;'
            f'padding:12px 16px;margin-bottom:16px">'
            f'<div style="font-size:11px;font-weight:700;letter-spacing:.5px;color:#8a8a8a;'
            f'text-transform:uppercase">{_esc(_field_label(template,"contingency"))}</div>'
            f'<div style="font-size:14px;color:#1f2937;margin-top:4px">{_esc(contingency)}</div></div>'
        )

    # Contacts
    contact_cells = ""
    for k in ("crisis_lead", "vendors", "siam_contact"):
        v = _val(values, k)
        if v:
            contact_cells += _box(_field_label(template, k), v, "#515151")
            contact_cells += '<td style="width:14px"></td>'
    if contact_cells:
        parts.append(f'<table style="width:100%;margin-bottom:16px"><tr>{contact_cells}</tr></table>')

    # Footer
    footer = _esc(template.get("footer", ""))
    parts.append(
        f'<div style="border-top:1px solid #e8ecf0;margin-top:8px;padding-top:12px;'
        f'font-size:12px;color:#6b7280;line-height:1.5">{footer}</div>'
    )

    inner = "".join(parts)
    return (
        f'<!DOCTYPE html><html><head><meta charset="utf-8">'
        f'<meta http-equiv="Content-Type" content="text/html; charset=utf-8">'
        f'<meta name="viewport" content="width=device-width,initial-scale=1">'
        f'</head><body style="margin:0;background-color:#f9f9f9;font-family:{FONT_STACK}">'
        f'<div style="max-width:720px;margin:0 auto;background-color:#ffffff;padding:24px">{inner}</div>'
        f'</body></html>'
    )


@api_router.get("/drafts/{draft_id}/render-html")
async def render_draft_html(draft_id: str):
    doc = await db.drafts.find_one({"id": draft_id, "deleted_at": None}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Draft not found")
    template = await db.templates.find_one({"id": doc["template_id"]}, {"_id": 0})
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    subject = _val(doc.get("values", {}), "subject") or doc.get("template_name", "")
    html = _render_html(template, doc.get("values", {}))
    to = await _compute_recipients(doc)
    return {"subject": subject, "html": html, "to": to}


# ---------------------------------------------------------------------------
# Excel import
# ---------------------------------------------------------------------------
@api_router.post("/import/excel")
async def import_excel(file: UploadFile = File(...)):
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Please upload an .xlsx file")
    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read spreadsheet: {exc}")

    added_products = 0
    added_processes = 0

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]

        def col_index(*names):
            for n in names:
                if n in header:
                    return header.index(n)
            return None

        product_col = col_index("product", "products", "business application",
                                "application", "affected service", "service",
                                "catalog product")
        platform_col = col_index("platform", "category", "group")
        process_col = col_index("business process", "process", "processes")

        for row in rows[1:]:
            if not row:
                continue
            if product_col is not None and product_col < len(row):
                name = row[product_col]
                if name and str(name).strip() and str(name).strip().lower() not in ("not in catalog", "none", "n/a"):
                    name = str(name).strip()
                    platform = "Imported"
                    if platform_col is not None and platform_col < len(row) and row[platform_col]:
                        platform = str(row[platform_col]).strip()
                    exists = await db.products.find_one({"name": name, "deleted_at": None})
                    if not exists:
                        await db.products.insert_one({
                            "id": new_id(), "name": name, "platform": platform,
                            "custom": True, "deleted_at": None, "created_at": now_iso(),
                        })
                        added_products += 1
            if process_col is not None and process_col < len(row):
                pname = row[process_col]
                if pname and str(pname).strip():
                    pname = str(pname).strip()
                    exists = await db.business_processes.find_one({"name": pname, "deleted_at": None})
                    if not exists:
                        await db.business_processes.insert_one({
                            "id": new_id(), "name": pname, "custom": True,
                            "deleted_at": None, "created_at": now_iso(),
                        })
                        added_processes += 1

    return {"added_products": added_products, "added_processes": added_processes}


@api_router.get("/contingency")
async def contingency_lookup(product: Optional[str] = None, process: Optional[str] = None):
    query: Dict[str, Any] = {}
    if product:
        query["product"] = product
    if process:
        query["process"] = process
    if not query:
        return []
    cursor = db.contingency_map.find(query, {"_id": 0})
    return await cursor.to_list(50)


@api_router.get("/contingency/export")
async def contingency_export():
    """Export the full Application → Business Process → Contingency plan mapping as .xlsx.

    Includes EVERY application (not only those with a contingency plan). Apps with no
    contingency row are still listed, with blank process/plan and 'No' in the flag column,
    so the application↔business-process mapping can be confirmed end-to-end.
    """
    products = await db.products.find({}, {"_id": 0}).to_list(2000)
    cont = await db.contingency_map.find({}, {"_id": 0}).to_list(2000)

    by_prod: Dict[str, List[Dict[str, Any]]] = {}
    for c in cont:
        by_prod.setdefault(c.get("product", ""), []).append(c)

    out: List[List[str]] = []
    seen = set()
    for p in sorted(products, key=lambda x: (str(x.get("platform", "")).lower(), str(x.get("name", "")).lower())):
        name = p.get("name", "")
        platform = p.get("platform", "")
        seen.add(name)
        matches = by_prod.get(name, [])
        if matches:
            for c in matches:
                out.append([platform, name, c.get("process", ""), c.get("failed_system", ""), c.get("owner", ""), "Yes"])
        else:
            out.append([platform, name, "", "", "", "No"])
    # contingency rows whose application is not in the product catalog — don't lose them
    for prod_name, matches in sorted(by_prod.items(), key=lambda x: x[0].lower()):
        if prod_name in seen:
            continue
        for c in matches:
            out.append(["", prod_name, c.get("process", ""), c.get("failed_system", ""), c.get("owner", ""), "Yes"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contingency Mapping"
    headers = [
        "Platform",
        "Business Application",
        "Business Process",
        "Contingency Plan (Failed System)",
        "Owner",
        "Has Contingency Plan",
    ]
    ws.append(headers)

    header_fill = openpyxl.styles.PatternFill("solid", fgColor="E61A27")
    header_font = openpyxl.styles.Font(bold=True, color="FFFFFF")
    thin = openpyxl.styles.Side(style="thin", color="D9D9D9")
    border = openpyxl.styles.Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = openpyxl.styles.Alignment(horizontal="left", vertical="center")
        cell.border = border

    muted_fill = openpyxl.styles.PatternFill("solid", fgColor="FBE9EA")
    for r in out:
        ws.append(r)
        if r[5] == "No":
            ws.cell(row=ws.max_row, column=6).fill = muted_fill

    widths = [26, 38, 46, 30, 22, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = openpyxl.styles.Alignment(vertical="top", wrap_text=True)
            cell.border = border
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"contingency_mapping_{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@api_router.get("/")
async def root():
    return {"message": "OpsComm API"}


@api_router.post("/outlook/drafts/{draft_id}")
async def outlook_create_draft(draft_id: str, authorization: str = Header(default="")):
    """Create the fully-formatted email as a DRAFT in the shared mailbox so the
    user opens Outlook desktop and presses Send (no copy/paste, no format loss)."""
    if not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Sign in to Microsoft first")
    doc = await db.drafts.find_one({"id": draft_id, "deleted_at": None}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Draft not found")
    template = await db.templates.find_one({"id": doc["template_id"]}, {"_id": 0})
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    subject = _val(doc.get("values", {}), "subject") or doc.get("template_name", "")
    html = _render_html(template, doc.get("values", {}))
    to = await _compute_recipients(doc)
    if not to:
        raise HTTPException(status_code=400, detail="No recipients — select at least one country first")
    result = await _create_shared_draft(authorization[7:], subject, to, html)
    return {"subject": subject, "to": to, **result}


register_outlook_routes(api_router, db)

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)
